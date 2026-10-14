#!/usr/bin/env python3
"""Importa a planilha de despesas e custos (uma por CNPJ) para o Financeiro da empresa correspondente.

Layout esperado (uma aba por mês):
  linha 1: A=Data, B=Fornecedor, C=NF, D=Boleto, E=Valor (R$), F=custo ou despesa, G..AK = um dia do mês por coluna
           (a data do pagamento é a data da coluna), AL=total.
  linhas de título: A=data de emissão, B=fornecedor, C=NF, D=boleto (S/N/$), E=valor do título,
           F=CUSTO/DESPESA, e o valor pago em cada dia na coluna daquele dia.
  linhas de folha: A=nome, B=função, F=CUSTO/DESPESA, valores pagos nos dias.

Cada valor numa coluna de dia vira UMA parcela de conta a pagar (Paga se a data já passou, Aberta se for futura)
e UM lançamento em Entradas e Saídas (Realizado ou Previsto), classificado em categoria/DRE.
Rodar de novo apaga o que esta importação criou antes (marcador "Importacao planilha ND" nas observações) e recria.

Uso:  python importar-planilha-despesas.py --dry-run
      ND_EMAIL=... ND_PASSWORD=... python importar-planilha-despesas.py --apply
"""
import argparse, urllib.parse, collections, datetime, glob, json, os, re, sys, urllib.request, urllib.error, uuid, http.cookiejar

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import nd_classificacao as nd

NAMES = {}
MARKER = "Importacao planilha ND"
NS = uuid.UUID("6f1c2d4e-9a3b-4c5d-8e7f-0a1b2c3d4e5f")
BASE = os.environ.get("ND_BASE_URL", "http://localhost:3000")
TODAY = datetime.date.today()
FILE_RE = re.compile(r"CNPJ\s*(\d{8})")


def uid(*parts):
    return str(uuid.uuid5(NS, "|".join(str(p) for p in parts)))


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def money(value):
    return round(float(value), 2)


def klass(value):
    text = clean(value).upper()
    return text if text in ("CUSTO", "DESPESA") else ""


def doc_number(value):
    if value is None or clean(value) == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean(value)


def method(boleto):
    text = clean(boleto)
    if text.upper() == "S":
        return "Boleto"
    if text == "$":
        return "Dinheiro"
    return ""


def read_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    payments, sheets = [], []
    for ws in wb.worksheets:
        days = {c.column: c.value.date() for c in ws[1] if isinstance(c.value, datetime.datetime)}
        if not days:
            continue
        total_row = None
        sheet_sum = 0.0
        for r in range(2, ws.max_row + 1):
            a, b, c, d, e, f = (ws.cell(r, i).value for i in range(1, 7))
            if a is None and b is None and f is None:
                if any(isinstance(ws.cell(r, col).value, (int, float)) for col in days):
                    total_row = r
                continue
            cells = [(col, ws.cell(r, col).value) for col in days if isinstance(ws.cell(r, col).value, (int, float)) and ws.cell(r, col).value]
            title = isinstance(a, datetime.datetime)
            for col, value in cells:
                sheet_sum += value
                payments.append(dict(sheet=ws.title.strip(), row=r, col=col, date=days[col], amount=money(value), is_title=title,
                                     issue=a.date() if title else None, party=clean(b) if title else clean(a), role="" if title else clean(b),
                                     nf=doc_number(c) if title else "", boleto=clean(d) if title else "", total=money(e) if isinstance(e, (int, float)) else 0.0,
                                     klass=klass(f), cell_total=ws.cell(r, 38).value))
        declared = ws.cell(total_row, 38).value if total_row else None
        sheets.append(dict(sheet=ws.title.strip(), imported=round(sheet_sum, 2), declared=declared if isinstance(declared, (int, float)) else None))
    return payments, sheets


def build_names(all_payments):
    """grafia padrão de cada nome: junta variações só de pontuação/espaço/acento (W. GARBOS = W GARBOS = W.GARBOS)."""
    groups = collections.defaultdict(collections.Counter)
    for p in all_payments:
        if p["is_title"]:
            groups[re.sub(r"[^A-Z0-9]", "", nd.norm(p["party"]))][p["party"]] += 1
    best = {}
    for key, spellings in groups.items():
        # prefere a grafia com acento, depois a mais frequente
        best[key] = sorted(spellings.items(), key=lambda kv: (-sum(ord(c) > 127 for c in kv[0]), -kv[1]))[0][0]
    return {nd.norm(p["party"]): best[re.sub(r"[^A-Z0-9]", "", nd.norm(p["party"]))] for p in all_payments if p["is_title"]}


def classify(p, names):
    """-> dict(canonical, supplier, supply, category, cc, dre): categoria financeira, fornecedor e conta do DRE de uma parcela."""
    if p["is_title"]:
        info = nd.classify_party(p["party"], p["klass"], names)
    else:
        info = dict(canonical=p["party"], supplier=False, supply=None, category=nd.PAYROLL.get(p["klass"], "Sem classificacao"), rule="(folha)")
    if p["klass"] == "CUSTO":
        dre, cc = nd.CUSTOS, "Producao"
    else:
        dre = nd.FIN_CATEGORIES[info["category"]]
        dre = "Administrativo" if dre == nd.CUSTOS else dre
        cc = "Administrativo" if p["klass"] == "DESPESA" else ""
    return {**info, "cc": cc, "dre": dre}


class Api:
    def __init__(self, email, password):
        self.jar = http.cookiejar.CookieJar()
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(self.jar))
        self.call("/api/auth/login", "POST", {"email": email, "password": password})

    def call(self, path, method="GET", body=None):
        req = urllib.request.Request(BASE + path, method=method, data=json.dumps(body).encode() if body is not None else None, headers={"Content-Type": "application/json"})
        try:
            with self.opener.open(req, timeout=300) as r:
                return json.loads(r.read().decode("utf-8") or "null")
        except urllib.error.HTTPError as err:
            raise SystemExit(f"{method} {path}: HTTP {err.code} {err.read().decode('utf-8', 'replace')[:300]}")


def rest(env, path, method="GET", body=None, prefer=None):
    headers = {"apikey": env["SUPABASE_SERVICE_ROLE_KEY"], "Authorization": "Bearer " + env["SUPABASE_SERVICE_ROLE_KEY"], "Content-Type": "application/json"}
    if prefer:
        headers["Prefer"] = prefer
    req = urllib.request.Request(env["NEXT_PUBLIC_SUPABASE_URL"] + "/rest/v1/" + path, method=method, data=json.dumps(body).encode() if body is not None else None, headers=headers)
    with urllib.request.urlopen(req, timeout=120) as r:
        return json.loads(r.read().decode("utf-8") or "null")


def write_review(plans, directory):
    """planilha de conferencia: nome na planilha -> fornecedor padronizado -> categorias (para revisar e corrigir as regras)."""
    rows = collections.OrderedDict()
    for _, company, payments, _ in plans:
        for p in payments:
            info = classify(p, NAMES)
            key = (p["party"], info["canonical"], p["is_title"])
            row = rows.setdefault(key, dict(info=info, total=0.0, count=0, companies=set(), klass=collections.Counter()))
            row["total"] += p["amount"]
            row["count"] += 1
            row["companies"].add(company["name"])
            row["klass"][p["klass"] or "-"] += 1
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "De-para"
    ws.append(["Nome na planilha", "Tipo", "Fornecedor padronizado", "Cria cadastro em Fornecedores?", "Categoria de fornecimento", "Categoria financeira", "Conta do DRE", "Classe na planilha", "Parcelas", "Total (R$)", "Empresas", "Regra"])
    for (name, canonical, title), row in sorted(rows.items(), key=lambda kv: -kv[1]["total"]):
        info = row["info"]
        ws.append([name, "Titulo/fornecedor" if title else "Folha de pagamento", canonical, "Sim" if info["supplier"] and title else "Nao", info.get("supply") or "", info["category"],
                   info["dre"], row["klass"].most_common(1)[0][0], row["count"], round(row["total"], 2), ", ".join(sorted(row["companies"])), info["rule"]])
    cat = wb.create_sheet("Categorias")
    cat.append(["Categoria financeira", "Conta do DRE (padrao)", "Parcelas", "Total (R$)"])
    totals = collections.defaultdict(lambda: [0, 0.0])
    for row in rows.values():
        totals[row["info"]["category"]][0] += row["count"]
        totals[row["info"]["category"]][1] += row["total"]
    for name, (count, total) in sorted(totals.items(), key=lambda kv: -kv[1][1]):
        cat.append([name, nd.FIN_CATEGORIES[name], count, round(total, 2)])
    for sheet in (ws, cat):
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(48, max(12, max(len(str(c.value or "")) for c in col[:60]) + 2))
    out = os.path.join(directory, "De-para fornecedores e categorias.xlsx")
    wb.save(out)
    print("\nPlanilha de conferencia gravada: " + os.path.abspath(out))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--dir", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
    parser.add_argument("--env", default=r"C:\carretinhas-dev\.env.local")
    args = parser.parse_args()
    files = sorted(glob.glob(os.path.join(args.dir, "Planilha_ND*.xlsx")))
    if not files:
        raise SystemExit("Nenhuma planilha Planilha_ND*.xlsx encontrada.")
    env = dict(l.strip().split("=", 1) for l in open(args.env, encoding="utf-8") if "=" in l and not l.startswith("#"))
    companies = {re.sub(r"\D", "", c["cnpj"] or "")[:8]: c for c in rest(env, "system_companies?select=id,name,cnpj")}

    plans = []
    for path in files:
        root = FILE_RE.search(os.path.basename(path)).group(1)
        company = companies.get(root)
        if not company:
            raise SystemExit(f"Empresa com CNPJ {root} nao cadastrada.")
        payments, sheets = read_workbook(path)
        plans.append((path, company, payments, sheets))
        print(f"\n=== {os.path.basename(path)} -> {company['name']} ({company['cnpj']})")
        print(f"{'aba':10} {'importado':>14} {'total da planilha':>18} {'diferenca':>12}")
        for s in sheets:
            diff = "" if s["declared"] is None else f"{s['imported'] - s['declared']:,.2f}"
            print(f"{s['sheet']:10} {s['imported']:>14,.2f} {('' if s['declared'] is None else format(s['declared'], ',.2f')):>18} {diff:>12}")
        paid = [p for p in payments if p["date"] <= TODAY]
        print(f"parcelas: {len(payments)} (pagas ate hoje: {len(paid)}, futuras: {len(payments) - len(paid)}) | soma {sum(p['amount'] for p in payments):,.2f}")
        for k, v in collections.Counter(p["klass"] or "(sem classe)" for p in payments).items():
            print(f"  {k}: {v} parcelas, {sum(p['amount'] for p in payments if (p['klass'] or '(sem classe)') == k):,.2f}")

    NAMES.update(build_names([p for _, _, payments, _ in plans for p in payments]))
    write_review(plans, args.dir)
    if not args.apply:
        print("\n(simulacao: nada foi gravado; use --apply)")
        return

    email, password = os.environ["ND_EMAIL"], os.environ["ND_PASSWORD"]
    api = Api(email, password)
    # fornecedores (cadastro único, vale para todas as empresas)
    suppliers = {}
    for _, _, payments, _ in plans:
        for p in payments:
            info = classify(p, NAMES)
            if info["supplier"] and p["is_title"]:
                suppliers.setdefault(info["canonical"], info["supply"])
    rows = [dict(id=uid("supplier", nd.norm(name)), name=name, document="21.161.569/0001-09" if name.startswith("FORT COM.") else "", person_type="PJ", category=supply,
                 status="Ativo", country="Brasil", created_by="Importacao planilha", commercial_notes="Criado pela importacao da planilha de despesas (sem CNPJ). Complete os dados.")
            for name, supply in sorted(suppliers.items())]
    for i in range(0, len(rows), 40):
        rest(env, "suppliers?on_conflict=id", "POST", rows[i:i + 40], prefer="resolution=ignore-duplicates,return=minimal")
    print(f"{len(rows)} fornecedores cadastrados/atualizados.")
    supplier_id = {name: uid("supplier", nd.norm(name)) for name in suppliers}
    for path, company, payments, _ in plans:
        api.call("/api/auth/companies", "POST", {"companyId": company["id"]})
        # apaga a importacao anterior desta empresa (marcador nas observacoes)
        for table in ("accounts_payable", "financial_transactions"):
            while True:  # em lotes: o banco tem gatilho de auditoria por linha e estoura o tempo em lotes grandes
                ids = [row["id"] for row in rest(env, f"{table}?select=id&system_company_id=eq.{company['id']}&notes=like.*{urllib.parse.quote(MARKER)}*&limit=200")]
                if not ids:
                    break
                rest(env, f"{table}?id=in.({','.join(ids)})", "DELETE")
        for old in ("Custos - Fornecedores", "Despesas - Fornecedores"):  # categorias da primeira versao da importacao
            rest(env, f"financial_categories?system_company_id=eq.{company['id']}&group=eq.Importacao%20planilha&name=eq.{urllib.parse.quote(old)}", "DELETE")
        state = api.call("/api/financial-state?refresh=1")["state"]
        cc = {c["name"]: c["id"] for c in state["costCenters"]}
        dre = {d["name"]: d["id"] for d in state["dreAccounts"]}
        cats = {c["name"]: c for c in state["categories"]}
        now = datetime.datetime.utcnow().isoformat() + "Z"
        used = {classify(p, NAMES)["category"] for p in payments}
        for name in sorted(used):
            if name not in cats:
                cat = dict(id=uid("cat", company["id"], name), name=name, group="Importacao planilha", type="saida", dreAccountId=dre[nd.FIN_CATEGORIES[name]], status="Ativo", createdAt=now, updatedAt=now, code="")
                state["categories"].append(cat)
                cats[name] = cat
        group_counts = collections.Counter()
        group_seen = collections.Counter()
        for p in payments:
            if p["is_title"]:
                group_counts[(p["party"], p["nf"], p["issue"], p["total"])] += 1
        new_payables, new_transactions = [], []
        for p in payments:
            info = classify(p, NAMES)
            cat_name, cc_name = info["category"], info["cc"]
            cat = cats[cat_name]
            party = info["canonical"] if info["supplier"] and p["is_title"] else p["party"]
            dre_id = dre[info["dre"]]
            paid = p["date"] <= TODAY
            key = (p["party"], p["nf"], p["issue"], p["total"])
            title = p["is_title"]
            count = group_counts[key] if title else 1
            group_seen[key] += 1
            number = group_seen[key] if title else 1
            description = (f"{party} - NF {p['nf']}" if p["nf"] else party) if title else f"Folha - {p['party']}{' (' + p['role'] + ')' if p['role'] else ''}"
            notes = "\n".join(filter(None, [f"{MARKER} | {p['sheet']} | linha {p['row']}",
                                            f"Emissao: {p['issue'].strftime('%d/%m/%Y')} | Valor do titulo: R$ {p['total']:,.2f} | Boleto: {p['boleto'] or '-'}" if title else f"Funcao: {p['role'] or '-'}",
                                            f"Classe na planilha: {p['klass'] or 'nao informada'}" + (f" | Nome na planilha: {p['party']}" if title and party != p["party"] else "")]))
            competence = min(p["issue"], p["date"]) if title and p["issue"] else p["date"]
            ap_id, ft_id = uid("ap", company["id"], p["sheet"], p["row"], p["col"]), uid("ft", company["id"], p["sheet"], p["row"], p["col"])
            new_payables.append(dict(id=ap_id, supplierId=supplier_id.get(info["canonical"], "") if info["supplier"] and title else "", supplierName=party, description=description, categoryId=cat["id"], subcategoryId="", costCenterId=cc.get(cc_name, ""), dreAccountId=dre_id,
                                     serviceOrderId="", providerId="", vehicleId="", competenceDate=competence.isoformat(), dueDate=p["date"].isoformat(), paymentDate=p["date"].isoformat() if paid else "",
                                     expectedAmount=p["amount"], paidAmount=p["amount"] if paid else 0, paymentMethod=method(p["boleto"]), bankAccountId="", creditCardId="", creditCardInvoiceId="",
                                     status="Paga" if paid else "Aberta", origin="Importacao planilha", notes=notes, attachmentName=os.path.basename(path), transactionId=ft_id, documentNumber=p["nf"],
                                     installmentNumber=number if count > 1 else None, installmentCount=count if count > 1 else None, installmentGroupId=uid("grp", company["id"], *map(str, key)) if title and count > 1 else "",
                                     paymentConditionId="", interestAmount=0, fineAmount=0, discountAmount=0, sourceType="Manual", sourceReference="", createdBy="Importacao planilha", updatedBy="", createdAt=now, updatedAt=now))
            new_transactions.append(dict(id=ft_id, type="Saida", description=description, categoryId=cat["id"], subcategoryId="", costCenterId=cc.get(cc_name, ""), dreAccountId=dre_id, clientId="", serviceOrderId="",
                                         providerId="", vehicleId="", supplierName=party, competenceDate=p["date"].isoformat(), dueDate=p["date"].isoformat(), realizedDate=p["date"].isoformat() if paid else "",
                                         expectedAmount=p["amount"], realizedAmount=p["amount"] if paid else 0, paymentMethod=method(p["boleto"]), bankAccountId="", creditCardId="", creditCardInvoiceId="",
                                         status="Realizado" if paid else "Previsto", origin="Conta a pagar" if paid else "Importacao planilha", notes=notes, attachmentName=os.path.basename(path), createdAt=now, updatedAt=now))
        batch = 120
        for i in range(0, len(new_payables), batch):
            state["accountsPayable"] = new_payables[i:i + batch]
            state["transactions"] = new_transactions[i:i + batch]
            api.call("/api/financial-state", "PUT", {"state": state})
            print(f"  {company['name']}: {min(i + batch, len(new_payables))}/{len(new_payables)}", flush=True)
        print(f"{company['name']}: {len(new_payables)} contas a pagar e {len(new_transactions)} lancamentos gravados.")


if __name__ == "__main__":
    main()
